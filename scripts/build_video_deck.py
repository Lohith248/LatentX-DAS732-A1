"""18-slide LatentX deck. Six slides each. No on-slide script cues."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(r"C:\Users\jagat\Desktop\Sem 5\DV")
IMG = ROOT / "images"
CROP = ROOT / "images" / "deck_crops"
OUTS = [
    ROOT / "LatentX_DAS732_A1_Video_Deck_v2.pptx",
    Path(r"C:\Users\jagat\Desktop\LatentX_DAS732_A1_Video_Deck_v2.pptx"),
]

W, H = Inches(13.333), Inches(7.5)
BG = RGBColor(0x0E, 0x0E, 0x10)
CARD = RGBColor(0x17, 0x17, 0x1B)
LINE = RGBColor(0x2A, 0x2A, 0x30)
RED = RGBColor(0xE1, 0x06, 0x2C)
INK = RGBColor(0xF4, 0xF1, 0xEA)
MUTED = RGBColor(0xB8, 0xB3, 0xA8)
DIM = RGBColor(0x7A, 0x75, 0x6C)

CROPS = {
    "Fig1.png": (0, 56, 0, 0),
    "Fig2.png": (0, 92, 0, 52),
    "Fig6.png": (0, 46, 0, 0),
    "Fig7.png": (0, 82, 0, 0),
    "Fig9.png": (0, 92, 0, 0),
    "Fig11.png": (0, 48, 0, 52),
    "Fig12.png": (0, 40, 0, 68),
    "Fig16.png": (0, 82, 0, 42),
    "Fig18.png": (0, 42, 0, 38),
    "Fig19.png": (0, 36, 0, 0),
    "Fig21.png": (0, 40, 0, 0),
    "Fig28.png": (0, 38, 0, 0),
}


def _set_run(run, text, size, color, bold=False, italic=False, font="Calibri"):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", font)


def _tf(shape, valign=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.10)
    tf.margin_right = Inches(0.10)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign],
        )
    except Exception:
        pass
    return tf


def add_text(slide, l, t, w, h, lines, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = _tf(box, valign)
    for i, spec in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(spec.get("after", 6))
        p.space_before = Pt(spec.get("before", 0))
        p.line_spacing = spec.get("lsp", 1.18)
        run = p.add_run()
        _set_run(
            run,
            spec["text"],
            spec.get("size", 16),
            spec.get("color", INK),
            spec.get("bold", False),
            spec.get("italic", False),
        )
    return box


def rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def pill(slide, l, t, w, h, fill=None, line=RED):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = 0.5
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1.25)
    return sh


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def page_num(slide, n, total=18):
    add_text(
        slide,
        Inches(12.15),
        Inches(7.18),
        Inches(0.95),
        Inches(0.24),
        [{"text": f"{n}  /  {total}", "size": 12, "color": DIM, "align": PP_ALIGN.RIGHT, "after": 0}],
    )


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def crop_figs():
    CROP.mkdir(parents=True, exist_ok=True)
    out = {}
    for name, (left, top, right, bottom) in CROPS.items():
        src = Image.open(IMG / name)
        w, h = src.size
        dest = CROP / name
        src.crop((left, top, w - right, h - bottom)).save(dest)
        out[name] = dest
    return out


def fit_picture(slide, path, l, t, w, h):
    rect(slide, l, t, w, h, CARD, LINE)
    im = Image.open(path)
    iw, ih = im.size
    max_w = w - Inches(0.16)
    max_h = h - Inches(0.16)
    aspect = iw / ih
    if aspect > (max_w / max_h):
        pw, ph = max_w, max_w / aspect
    else:
        ph, pw = max_h, max_h * aspect
    slide.shapes.add_picture(str(path), l + (w - pw) / 2, t + (h - ph) / 2, pw, ph)


def title_block(slide, title, subtitle):
    add_text(
        slide,
        Inches(0.42),
        Inches(0.12),
        Inches(12.5),
        Inches(0.82),
        [{"text": title, "size": 24, "bold": True, "after": 0, "lsp": 1.08}],
    )
    add_text(
        slide,
        Inches(0.42),
        Inches(0.92),
        Inches(12.5),
        Inches(0.34),
        [{"text": subtitle, "size": 16, "color": MUTED, "after": 0}],
    )
    rect(slide, Inches(0.42), Inches(1.26), Inches(1.20), Inches(0.034), RED)


def evidence(prs, n, title, subtitle, fig, paras, caveat, note):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    title_block(slide, title, subtitle)
    fit_picture(slide, fig, Inches(0.42), Inches(1.46), Inches(8.22), Inches(5.47))
    rect(slide, Inches(8.82), Inches(1.46), Inches(4.08), Inches(5.47), CARD, LINE)
    lines = []
    for i, p in enumerate(paras):
        lines.append({"text": p, "size": 15, "after": 12 if i < len(paras) - 1 else 0, "lsp": 1.2})
    add_text(slide, Inches(8.94), Inches(1.58), Inches(3.84), Inches(4.15), lines)
    add_text(
        slide,
        Inches(8.94),
        Inches(5.78),
        Inches(3.84),
        Inches(0.95),
        [{"text": caveat, "size": 13, "italic": True, "color": MUTED, "after": 0, "lsp": 1.15}],
    )
    page_num(slide, n)
    notes(slide, note)
    return slide


def build():
    figs = crop_figs()
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 Lohith open + question
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    rect(s, Inches(0), Inches(0), Inches(0.09), H, RED)
    pill(s, Inches(0.50), Inches(0.32), Inches(2.45), Inches(0.38))
    add_text(
        s,
        Inches(0.50),
        Inches(0.33),
        Inches(2.45),
        Inches(0.36),
        [{"text": "TEAM LATENTX", "size": 14, "bold": True, "align": PP_ALIGN.CENTER, "after": 0}],
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        s,
        Inches(0.50),
        Inches(0.82),
        Inches(12.3),
        Inches(1.45),
        [
            {"text": "Reports rose through 2022", "size": 36, "bold": True, "after": 0, "lsp": 1.0},
            {"text": "and leveled in 2023", "size": 36, "bold": True, "after": 0, "lsp": 1.0},
        ],
    )
    add_text(
        s,
        Inches(0.50),
        Inches(2.32),
        Inches(12.3),
        Inches(0.40),
        [
            {
                "text": "Official LAPD legacy reports, 2020–2024. Several visible changes are recording artefacts.",
                "size": 16,
                "color": MUTED,
                "after": 0,
            }
        ],
    )
    for i, (name, roll) in enumerate(
        (("Lohith P", "BT2024248"), ("Sri Charan", "BT2024143"), ("R Anish Reddy", "BT2024228"))
    ):
        x = Inches(0.50) + Inches(3.40) * i
        add_text(
            s,
            x,
            Inches(2.78),
            Inches(3.20),
            Inches(0.70),
            [
                {"text": name, "size": 16, "bold": True, "after": 2},
                {"text": roll, "size": 13, "color": MUTED, "after": 0},
            ],
        )
    rect(s, Inches(0.50), Inches(3.60), Inches(12.32), Inches(2.85), CARD, LINE)
    add_text(
        s,
        Inches(0.70),
        Inches(3.72),
        Inches(11.95),
        Inches(0.32),
        [{"text": "THE QUESTION", "size": 12, "bold": True, "color": RED, "after": 0}],
    )
    add_text(
        s,
        Inches(0.70),
        Inches(4.08),
        Inches(11.95),
        Inches(2.15),
        [
            {
                "text": "Between 2020 and 2024, when, where, and what kind of crime did Los Angeles residents report to the LAPD, and which visible changes are recording artefacts rather than patterns?",
                "size": 20,
                "lsp": 1.25,
                "after": 0,
            }
        ],
    )
    page_num(s, 1)
    notes(
        s,
        "Hi everyone, we are Team LatentX. I am Lohith. With me are Sri Charan and R Anish Reddy. We spent this project on one official LAPD file, 2020 to 2024, and on a pretty stubborn question: when you look at reported crime in Los Angeles, which swings are real patterns in the reports, and which ones are just how the records system was filled in? I will start with time and with how we cleaned the file. Sri Charan takes place. Anish takes crime type and then he closes.",
    )

    # 2 people + source
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title_block(
        s,
        "Three people, one official freeze",
        "Each person owns a full task. One row in this file is one report, not a crime rate.",
    )
    trio = [
        ("Lohith P", "Time and the records system. I check the trend, the 2024 coverage break, and the noon and day-one defaults."),
        ("Sri Charan", "Place. He separates how many reports a division has from what mix of crimes it has, then ranks the districts."),
        ("R Anish Reddy", "Crime type, then the close. He splits volume from growth and looks hard at identity theft."),
    ]
    for i, (name, body) in enumerate(trio):
        x = Inches(0.42) + Inches(4.22) * i
        rect(s, x, Inches(1.50), Inches(4.04), Inches(2.55), CARD, LINE)
        add_text(s, x + Inches(0.16), Inches(1.62), Inches(3.72), Inches(0.40), [{"text": name, "size": 18, "bold": True, "after": 0}])
        add_text(s, x + Inches(0.16), Inches(2.08), Inches(3.72), Inches(1.80), [{"text": body, "size": 14, "lsp": 1.2, "after": 0}])
    facts = [
        ("1,004,894", "reports in the frozen official table"),
        ("DR_NO", "one report, one row"),
        ("2020–2023", "the only full comparable years"),
        ("7 Mar 2024", "LAPD left this table for NIBRS"),
    ]
    for i, (num, lab) in enumerate(facts):
        x = Inches(0.42) + Inches(3.22) * i
        rect(s, x, Inches(4.24), Inches(3.08), Inches(2.15), CARD, LINE)
        add_text(s, x + Inches(0.14), Inches(4.36), Inches(2.80), Inches(0.55), [{"text": num, "size": 20, "bold": True, "after": 0}])
        add_text(s, x + Inches(0.14), Inches(4.95), Inches(2.80), Inches(1.20), [{"text": lab, "size": 13, "color": MUTED, "lsp": 1.15, "after": 0}])
    add_text(
        s,
        Inches(0.42),
        Inches(6.50),
        Inches(12.4),
        Inches(0.45),
        [
            {
                "text": "Assigned Kaggle link returned 403. Staff said use the better source. We kept this official freeze, CC0, table 2nrs-mtv8, for A1, A2, and A3.",
                "size": 14,
                "color": MUTED,
                "after": 0,
            }
        ],
    )
    page_num(s, 2)
    notes(
        s,
        "A few facts before any chart. The Kaggle link we were assigned came back 403. Staff said use whichever source is better, so we froze the official LAPD table: one million four thousand eight hundred ninety-four reports. Each row is one DR_NO, one report, not a person, not a rate. 2020 through 2023 are complete years. On 7 March 2024 the department left this legacy system. I read time against that window. Sri Charan reads place. Anish reads type, and he will bring us back at the end.",
    )

    # 3 preprocess
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title_block(
        s,
        "What we did to one million rows before Tableau",
        "Same freeze for every chart. Tableau built by hand. NIBRS was not merged.",
    )
    steps = [
        ("1", "Lock the official file", "1,004,894 rows. We do not add later dumps or the NIBRS tables. Those have a different grain and they do not finish 2024."),
        ("2", "Parse dates and delay", "Occurrence date, report date, and the days between them. That delay is how we later see identity theft arriving late."),
        ("3", "Flag recording defaults", "Exact noon: 35,198 rows. First of the month: 46,546 rows. Those flags sit on the time charts, not in a footnote."),
        ("4", "Clean coordinates and victims", "2,240 rows sit at 0,0 and get nulled. Victim age and sex are cleaned. Descent is not used."),
        ("5", "Derive families, then split views", "Crime families come from the description. Totals use every relevant row. Point maps use 200,000 rows, seed 732."),
        ("6", "Compare 2020–2023", "2024 stays on the time slides as a coverage break. Place and type charts stop at 2023."),
    ]
    for i, (num, head, body) in enumerate(steps):
        col, row = i % 3, i // 3
        x = Inches(0.42) + Inches(4.22) * col
        y = Inches(1.50) + Inches(2.55) * row
        rect(s, x, y, Inches(4.04), Inches(2.40), CARD, LINE)
        add_text(s, x + Inches(0.16), y + Inches(0.10), Inches(0.40), Inches(0.36), [{"text": num, "size": 18, "bold": True, "color": RED, "after": 0}])
        add_text(s, x + Inches(0.52), y + Inches(0.12), Inches(3.35), Inches(0.36), [{"text": head, "size": 16, "bold": True, "after": 0}])
        add_text(s, x + Inches(0.16), y + Inches(0.55), Inches(3.72), Inches(1.70), [{"text": body, "size": 13, "lsp": 1.18, "after": 0}])
    page_num(s, 3)
    notes(
        s,
        "Here is the prep, in the order we actually did it. We locked that official freeze and we did not stitch in NIBRS, because NIBRS is one row per offense and it does not complete 2024. We parsed the occurrence date, the report date, and the delay between them. We flagged two defaults we already knew would lie: exact noon, 35,198 rows, and the first of the month, 46,546. We threw out 2,240 coordinates sitting on 0,0. We grouped descriptions into crime families. Totals always use the full relevant rows. The point maps use a 200,000-row sample so Tableau does not die. And from here on, if I compare years, I mean 2020 to 2023. 2024 is only allowed on the time charts, as a broken series.",
    )

    evidence(
        prs,
        4,
        "Reports rose through 2022 and leveled in 2023",
        "X-axis is occurrence date. Y-axis is reports that day. Grey is daily. Navy is a 28-day mean.",
        figs["Fig1.png"],
        [
            "Left to right is time, 2020 into early 2025. Up is how many reports that day.",
            "The grey jitter is every day. The navy line is a 28-day smooth so the shape is readable.",
            "The series goes 199,847 in 2020, 235,259 in 2022, 232,345 in 2023, then falls off after 7 March 2024.",
        ],
        "That last drop is this table losing coverage, not a citywide crime year.",
        "Look at the axes first. Along the bottom is the occurrence date. Up the side is reports per day. The grey is supposed to look noisy; that is every single day. The navy line is a 28-day average, so you can see the shape without pretending the days are smooth. Spring 2020 dips, then the line climbs into 2022. The annual totals are 199,847 in 2020, 235,259 in 2022, and 232,345 in 2023, so the rise stops. Then this cliff after 7 March 2024. I am not going to call that a safer city. That is the department leaving this file. Next slide tests that claim month by month.",
    )
    evidence(
        prs,
        5,
        "2024 is missing coverage, not a safer city",
        "Each bar is one month of 2024. The label is that month as a percent of the 2020–2023 average for the same month.",
        figs["Fig2.png"],
        [
            "Y-axis is monthly reports. The navy line is the older same-month average.",
            "January to March 2024 still sit near that line: 103 percent, 101 percent, 91 percent.",
            "April is already 73 percent. December is 26 percent. The table is emptying after 7 March.",
        ],
        "Place and type charts therefore use 2020–2023 only.",
        "Same story, now by month. The height of each bar is 2024 reports. The number on the bar is 2024 as a percent of the average for that month in 2020 to 2023. January, February, March are basically on the old baseline: 103, 101, 91. Then the system change, and April is 73 percent. By December you are at 26 percent of a normal December. So if someone says crime collapsed in 2024, they are reading a file that stopped being fed. Sri Charan and Anish will only compare the complete years.",
    )
    evidence(
        prs,
        6,
        "Two spikes are recording defaults",
        "(a) hour of day. (b) day of month. (c) share of each family stamped exactly 12:00.",
        figs["Fig6.png"],
        [
            "Panel (a): hour 12 is the tower. That is 35,198 rows stamped 12:00, about 3.5 percent of the freeze.",
            "Panel (b): day 1 is the tallest date bar, 46,546 records. Empty dates get the first of the month.",
            "Panel (c): identity theft takes that noon stamp 11 percent of the time. The aggravated family only 1.5 percent.",
        ],
        "Hour 12 is a default timestamp, not a rush hour. I hand place to Sri Charan.",
        "These three panels are the artefacts I do not want you to carry into the next talks. Left chart, x is hour 0 to 23, y is reports. Hour 12 is the ugly tower: 35,198 rows, about 3.5 percent of the whole freeze, all stamped exactly noon. Middle chart, x is day of month. The first is tallest, 46,546, because missing dates get day one. Bottom chart asks who gets the noon stamp. Identity theft, 11 percent. Aggravated assault, 1.5 percent. So if a later chart lights up at noon, that is the form, not a lunch-hour crime wave. I stop on time here. Sri Charan takes place.",
    )

    evidence(
        prs,
        7,
        "Central leads volume, then 77th Street and Pacific",
        "21 divisions, 2020–2023. Darker navy is more reports. This is a count, not a rate.",
        figs["Fig9.png"],
        [
            "Geography is the map. Color is volume. There is no third variable hiding in the color.",
            "Central is darkest at 59,456 reports. Then 77th Street 54,981 and Pacific 51,351.",
            "Foothill is the lightest at 29,182. Darker means more recorded workload on hundred-blocks, not more danger.",
        ],
        "Not a population rate and not a prediction.",
        "Hello, I am Sri Charan. I am only talking about where reports sit. This is a choropleth of the 21 LAPD divisions for 2020 to 2023. The color scale on the bottom is report counts, 29,182 up to 59,456. Central is the darkest, 59,456. Then 77th Street, 54,981, and Pacific, 51,351. Foothill is pale at 29,182. Please do not turn that into a danger map. These are hundred-block addresses and raw counts. Next I check whether 2024 falling apart is a Central story or a citywide file story.",
    )
    evidence(
        prs,
        8,
        "Every division loses 2024 the same way",
        "Each panel is one division. X is month. Y is monthly reports. The late drop is 2024 in every panel.",
        figs["Fig11.png"],
        [
            "Do not read 21 tiny trends. Read that the cliff arrives in all 21 boxes at the same time.",
            "Central also shows the 2021–22 rise Lohith already counted. The 2024 break is not special to Central.",
            "If one neighborhood had simply become safer, you would not see the same cut in Foothill and in Central.",
        ],
        "This is the coverage loss, drawn in place. It is not a local crime drop.",
        "Twenty-one small charts, one per division. X is month, y is monthly reports. I am not going to walk each line. I want you to see one thing: in 2024 every panel falls off together. Central does have that earlier rise into 2022, which matches Lohith’s city total. But the 2024 cut is in Foothill, Harbor, West LA, all of them. So we do not say one area suddenly got safe. The whole legacy table thinned out. From here my mix and district charts stay on 2020 to 2023.",
    )
    evidence(
        prs,
        9,
        "The busiest divisions do not share one crime mix",
        "Each row is one division and adds to 100 percent. Color is that family’s share of the row.",
        figs["Fig12.png"],
        [
            "This is not the map again. Volume is gone. Every row is renormalized to 100 percent.",
            "Combined assault is 40.6 percent of 77th Street and 39.6 percent of Southeast. Aggravated alone is 24.6 percent in 77th.",
            "West LA is 21.1 percent other theft and 4.8 percent aggravated. Central is 18.0 percent theft from vehicle.",
        ],
        "Other and Other theft are different families. Do not collapse them.",
        "Now I take volume off the table. Each row is a division, and the row sums to 100 percent, so you are looking at mix. The x-axis is crime family. Darker teal is a bigger share of that division. 77th Street and Southeast are assault-heavy: about 40 percent combined assault, and 24.6 percent aggravated in 77th alone. Flip to West LA: 21 percent other theft, under 5 percent aggravated. Central, which won the volume map, is 18 percent theft from vehicle. So the busiest places are not busy with the same crime. Also, Other and Other theft are two families in this table. We never merge them.",
    )
    evidence(
        prs,
        10,
        "The noon stripe is a default. The real band is evening.",
        "Rows are divisions. Columns are hour 0 to 23. Color is that hour’s share of the division.",
        figs["Fig18.png"],
        [
            "The dark vertical stripe at hour 12 is Lohith’s noon default, up to 8.4 percent in Topanga.",
            "Once you ignore that stripe, the remaining dark cells sit between 17 and 20.",
            "Central peaks at hour 18, 6.0 percent. Hollenbeck at hour 20, 6.1 percent.",
        ],
        "Do not call hour 12 a rush hour. That column is a timestamp.",
        "Same 21 divisions, now by hour. Rows are divisions, columns are hour of day, color is the share of that division’s reports. You can see a dark stripe straight down hour 12. That is the default Lohith flagged, as high as 8.4 percent in Topanga. Mentally cover that column. What is left is evening, 17 to 20. Central’s own peak is 6 in the evening, 6 percent. Hollenbeck is 8 at night, 6.1 percent. So the map is not “crime happens at noon.” Noon is the form. The reports that have a real hour pile up after work.",
    )
    evidence(
        prs,
        11,
        "Concentration is real, and it has a long tail",
        "X is district rank, most reports to least. Blue bars are counts. The red line is the running share.",
        figs["Fig16.png"],
        [
            "1,207 reporting districts. The red line hits 80 percent at rank 623, labeled on the chart.",
            "Rank 623 is about half the districts, not a handful. The top bar is Central RD 0162, 4,672 reports.",
            "I hand type to Anish. Place can tell you where the paper sits. It cannot tell you what grew.",
        ],
        "Ranks are report volume, not risk.",
        "Last place chart. I ranked all 1,207 reporting districts from most reports to least. The blue mountain is those counts. The red line is the cumulative share, and it is labeled 80 percent at rank 623. That means you need 623 districts, about half of them, to cover 80 percent of reports. The single tallest bar is Central reporting district 0162, 4,672 reports, not 45,000. So yes, work concentrates, and no, a handful of blocks do not hold the city. I am done with place. Anish takes what was actually written on the report.",
    )

    # 12 place wrap
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title_block(
        s,
        "Place in three facts, then we change the question",
        "Volume, mix, and hour are not the same map. Anish picks up type from here.",
    )
    wrap = [
        ("Volume", "Central, 77th Street, and Pacific take the most reports. Foothill takes the fewest. That is workload, on hundred-blocks."),
        ("Mix", "77th Street and Southeast are more assault. West LA and parts of the Westside are more theft. High count is not one crime."),
        ("Hour", "The noon column is a default. After you set it aside, the remaining mass is 17:00 to 20:00."),
    ]
    for i, (h, b) in enumerate(wrap):
        x = Inches(0.42) + Inches(4.22) * i
        rect(s, x, Inches(1.55), Inches(4.04), Inches(4.55), CARD, LINE)
        add_text(s, x + Inches(0.22), Inches(1.80), Inches(3.60), Inches(0.50), [{"text": h, "size": 22, "bold": True, "color": RED, "after": 0}])
        add_text(s, x + Inches(0.22), Inches(2.45), Inches(3.60), Inches(3.30), [{"text": b, "size": 17, "lsp": 1.28, "after": 0}])
    page_num(s, 12)
    notes(
        s,
        "Let me leave place as three separate facts, because they get mashed together if I am not careful. One: Central, 77th, and Pacific take the paper. Two: they do not take the same crimes. 77th is assault-heavy, Central is more theft from cars. Three: noon is a recording stripe; the hours that behave like hours are evening. I am not saying why a division looks like that, and I am not ranking neighborhoods by danger. Anish, over to you for what grew.",
    )

    evidence(
        prs,
        13,
        "Most reports are filed the same day. Identity theft is not.",
        "Top: delay in days. Middle: each family’s delay mix. Bottom: occurrence year versus report year.",
        figs["Fig7.png"],
        [
            "48 percent, 482,015 reports, are filed the same day. Assault’s median delay is 0 days.",
            "Identity theft’s median is 7 days. That is the family I am about to show growing.",
            "The 94 rows with a 2025 report date are late filings of older incidents, not a 2025 crime year.",
        ],
        "Delay is a recording interval. It is not a motive.",
        "Hi, I am R Anish Reddy. I start from delay because it already points at the family that grew. The top panel is days between the incident and the filing. Almost half the file, 482,015 reports, is same day. In the middle panel, each bar is a crime family. Assault sits at a 0-day median. Identity theft sits at 7 days. The bottom grid is just occurrence year versus report year; those 94 filings stamped 2025 are late paperwork, not a new year of crime. So when I say identity theft grew, remember it also arrives in the file later than a street assault.",
    )
    evidence(
        prs,
        14,
        "Identity theft doubled by 2022. Assault stayed near its 2020 level.",
        "Every family is set to 100 in 2020. Y-axis is that index, not the raw count.",
        figs["Fig21.png"],
        [
            "Read this as growth from an unequal start. A family of 5,000 and a family of 50,000 both begin at 100.",
            "The red line is identity theft: 200.9 in 2022, still 139.2 in 2023.",
            "Other theft ends at 151.2. The assault lines stay near 108. Vandalism finishes below 100.",
        ],
        "An index cannot tell you which crime is the largest. That is the next slide.",
        "This chart is easy to misread, so start with the axes. X is year. Y is not reports. Y is an index: every family is 100 in 2020, so a small family and a huge family can be compared on growth. Identity theft, the red line, goes to 200.9 in 2022. That is a double. In 2023 it is still 139. Other theft finishes at 151. Look at the assault lines: they barely leave 108. Vandalism is actually under 100. So the rise Lohith counted is a theft-composition change. It is not the city becoming more violent in this file. Next slide puts the raw sizes back, because this chart hides who is still the biggest.",
    )
    evidence(
        prs,
        15,
        "Largest by count is still vehicle theft",
        "Tile size is 2020–2023 volume. Color is change from 2020 to 2023. Warm is up.",
        figs["Fig19.png"],
        [
            "The biggest tile is VEHICLE – STOLEN, 93,674 reports. Then simple battery, 69,421.",
            "THEFT OF IDENTITY is 56,216. It is not the largest tile. It is one of the warmer ones.",
            "Size answers “what fills the file.” Color answers “what changed.” Those are different questions.",
        ],
        "Do not read a warm tile as a cause or as guilt.",
        "Now ignore the index and look at area. This treemap is the file by description. Bigger tile, more reports. Vehicle stolen is still the giant, 93,674. Simple battery is 69,421. Identity theft is 56,216, so it is third, not first. The color is percent change from 2020 to 2023. Identity theft is warmer because it grew. Vehicle stolen is huge and not as warm. If you only watch the red line on the last slide you will think identity theft took over Los Angeles. It did not. It grew faster. The pile of paper is still vehicle theft.",
    )
    evidence(
        prs,
        16,
        "Identity theft peaks in late 2022 and is filed later",
        "Four coordinated views: month, recorded age, premise, and delay versus the city.",
        figs["Fig28.png"],
        [
            "Top left: monthly counts peak at 2,737 in December 2022, then the 2024 coverage cut arrives.",
            "Residence is 70.40 percent. Recorded ages pile up from 25 to 39.",
            "About 28 percent are filed after 31 days, against 5.9 percent citywide. We do not claim a cause.",
        ],
        "These are recorded fields, not every person affected.",
        "This is one family, four windows. Top left, x is month, y is identity-theft reports. You see the climb to 2,737 in December 2022, then the same 2024 coverage haircut Lohith showed for the whole file. Top right, recorded age: 25 to 39. Bottom left, premise: 70 percent residence. Bottom right, delay: the pink bars are identity theft, the dots are the city. About 28 percent of identity theft is filed after 31 days. Citywide that bucket is 5.9 percent. I am not going to invent a reason. I am saying the family that grew is also the family that is written down late, and mostly at a home address. That is as far as this table goes.",
    )

    # 17 synthesis Anish
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title_block(
        s,
        "The rise is a theft story. The 2024 cliff is a records break.",
        "What the three of us can stand behind, and what the file will not let us say.",
    )
    rect(s, Inches(0.42), Inches(1.48), Inches(6.08), Inches(2.85), CARD, LINE)
    rect(s, Inches(6.72), Inches(1.48), Inches(6.18), Inches(2.85), CARD, LINE)
    add_text(
        s,
        Inches(0.60),
        Inches(1.60),
        Inches(5.75),
        Inches(2.60),
        [
            {"text": "WHAT THE CHARTS SUPPORT", "size": 13, "bold": True, "color": RED, "after": 10},
            {
                "text": "Reports rose from 199,847 in 2020 to 235,259 in 2022 and leveled at 232,345 in 2023.",
                "size": 16,
                "after": 10,
                "lsp": 1.2,
            },
            {
                "text": "Volume sits in Central, 77th Street, and Pacific, but the mix is not one crime, and the growth is in theft, especially identity theft.",
                "size": 16,
                "after": 0,
                "lsp": 1.2,
            },
        ],
    )
    add_text(
        s,
        Inches(6.90),
        Inches(1.60),
        Inches(5.85),
        Inches(2.60),
        [
            {"text": "WHAT IS AN ARTEFACT", "size": 13, "bold": True, "color": RED, "after": 10},
            {
                "text": "After 7 March 2024 the legacy table loses coverage, 73 percent in April, 26 percent in December, in every division.",
                "size": 16,
                "after": 10,
                "lsp": 1.2,
            },
            {
                "text": "Exact noon, 35,198 reports, and the first of the month, 46,546 records, are defaults. Evening 17 to 20 is what remains.",
                "size": 16,
                "after": 0,
                "lsp": 1.2,
            },
        ],
    )
    rect(s, Inches(0.42), Inches(4.52), Inches(12.48), Inches(2.18), CARD, LINE)
    add_text(
        s,
        Inches(0.68),
        Inches(4.70),
        Inches(12.00),
        Inches(1.85),
        [
            {
                "text": "Reported crime in Los Angeles rose through 2022 and leveled in 2023, driven more by theft, especially identity theft, than by violence. It concentrates in Downtown and South LA. The 2024 coverage loss, the noon spike, and the first-of-month spike are recording artefacts, not crime trends.",
                "size": 17,
                "lsp": 1.25,
                "after": 0,
            }
        ],
    )
    page_num(s, 17)
    notes(
        s,
        "Let me put the three talks back into one sentence each, then one sentence for the whole project. Lohith: the reports rise into 2022, sit in 2023, and 2024 is a broken series. Sri Charan: the paper piles up Downtown and South LA, but mix and hour are different facts from volume, and noon is a default. Me: vehicle theft is still the biggest pile; identity theft is the one that doubled and is filed late. Together: reported crime in Los Angeles rose through 2022 and leveled in 2023, more a theft story than a violence story, concentrated in Downtown and South LA. The 2024 cliff, the noon spike, and the first-of-month spike are recording artefacts. We never turned these counts into rates, and we are not explaining why identity theft grew.",
    )

    # 18 close Anish
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    rect(s, Inches(0), Inches(0), Inches(0.09), H, RED)
    pill(s, Inches(0.52), Inches(1.45), Inches(2.45), Inches(0.40))
    add_text(
        s,
        Inches(0.52),
        Inches(1.46),
        Inches(2.45),
        Inches(0.38),
        [{"text": "TEAM LATENTX", "size": 14, "bold": True, "align": PP_ALIGN.CENTER, "after": 0}],
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        s,
        Inches(0.52),
        Inches(2.10),
        Inches(12.2),
        Inches(0.80),
        [{"text": "Thank you", "size": 48, "bold": True, "after": 0}],
    )
    add_text(
        s,
        Inches(0.52),
        Inches(3.00),
        Inches(12.2),
        Inches(0.70),
        [
            {
                "text": "Official freeze, preprocess script, three Tableau workbooks, and the IEEE report.",
                "size": 18,
                "color": MUTED,
                "after": 0,
            }
        ],
    )
    add_text(
        s,
        Inches(0.52),
        Inches(3.85),
        Inches(12.2),
        Inches(0.50),
        [{"text": "github.com/Lohith248/LatentX-DAS732-A1", "size": 24, "bold": True, "after": 0}],
    )
    add_text(
        s,
        Inches(0.52),
        Inches(4.55),
        Inches(12.2),
        Inches(0.45),
        [{"text": "Lohith P   ·   Sri Charan   ·   R Anish Reddy", "size": 18, "after": 0}],
    )
    add_text(
        s,
        Inches(0.52),
        Inches(6.85),
        Inches(10.5),
        Inches(0.30),
        [
            {
                "text": "DAS732 Data Visualization   ·   IIIT Bangalore   ·   T1 2026–27",
                "size": 14,
                "color": MUTED,
                "after": 0,
            }
        ],
    )
    page_num(s, 18)
    notes(
        s,
        "If you want the charts, the Python, or the three Tableau workbooks, they are in the GitHub repo on this slide. We are Team LatentX. I am Anish, that was Lohith and Sri Charan, and thank you for the time.",
    )

    for dest in OUTS:
        dest.parent.mkdir(parents=True, exist_ok=True)
        prs.save(dest)
        print("wrote", dest)


if __name__ == "__main__":
    build()
